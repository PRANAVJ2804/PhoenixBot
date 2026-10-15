import random
import time
import schedule
import gspread
import openpyxl
from datetime import datetime
from google.oauth2.service_account import Credentials
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from webdriver_manager.chrome import ChromeDriverManager

# ─────────────────────────────────────────────
#  QUOTES
# ─────────────────────────────────────────────
QUOTES = [
    "Every name on this list is a life being changed. Never forget that.",
    "The best time to close was yesterday. The next best time is right now.",
    "Champions show up every single day. This team shows up.",
    "One admission at a time. One dream fulfilled at a time.",
    "Your work today is someone's future tomorrow.",
    "Greatness is not a moment. It's a daily choice. You're making it.",
    "Behind every number is a student whose life just changed. Keep going.",
    "The scoreboard reflects the effort. Make the effort worth reflecting.",
    "Small consistent actions build extraordinary outcomes.",
    "This team doesn't just close admissions. It opens futures.",
    "Every admission is a promise kept to a student who believed in you.",
    "The numbers you build today will define the legacy of this team.",
    "Success is not given. It is earned — one admission at a time.",
    "You are not just selling education. You are changing destinies.",
    "The best counsellors don't wait for leads. They create opportunities.",
    "Effort compounds. Every application filed today matters tomorrow.",
    "Consistency is the bridge between goals and achievement.",
    "Great work doesn't announce itself. The numbers do. Keep going.",
    "Each admission on this list is a family's hope fulfilled.",
    "The grind is quiet. The results are loud. Keep grinding.",
    "You don't have to be the best today. Just better than yesterday.",
    "Progress is built in the moments no one sees. This team sees it.",
    "Every student admitted is a vote of confidence in this team.",
    "The scoreboard changes one effort at a time. Make yours count.",
    "Hard work in silence. Let success make the noise.",
    "This team doesn't count the days. It makes the days count.",
    "Impact is not measured in hours. It is measured in admissions.",
    "Stay focused. Stay consistent. The numbers will follow.",
    "Behind every great number is an even greater team effort.",
    "The best is not who works the hardest. It's who never stops.",
    "Dreams don't work unless you do. This team works.",
    "Every application filed is a seed planted for someone's future.",
    "Do the work today that your future self will thank you for.",
    "This is more than a job. This is a mission. Own it.",
    "The scoreboard tells a story. Make sure it's a great one.",
    "Momentum is built by people who refuse to stop. Be that person.",
    "Champions are not born. They are built in moments like this.",
    "Every closed admission is a chapter written in someone's success story.",
    "The difference between good and great is relentless consistency.",
    "Keep going. The numbers you build today echo for years.",
]

last_report_quote = None

def get_report_quote():
    global last_report_quote
    available = [q for q in QUOTES if q != last_report_quote]
    chosen = random.choice(available if available else QUOTES)
    last_report_quote = chosen
    return chosen

# ─────────────────────────────────────────────
#  CONFIG
# ─────────────────────────────────────────────
SHEET_ID          = "1EPI1eIECYm5DQCBarkmPlwscd6Ycn007Te1tysYIUeI"
TAB_APPS          = "All Universities applications"
TAB_ADM           = "Payment done Roll no pending"
CREDENTIALS_PATH  = r"C:\Users\USER\Desktop\whatsapp_bot\credentials.json"
CHROME_PROFILE    = r"C:\Users\USER\Desktop\whatsapp_bot\chrome_session_report"
COUNSELLORS_EXCEL = r"C:\Users\USER\Desktop\whatsapp_bot\Counsellors.xlsx"
WHATSAPP_GROUP_P  = "Team Prashanth"
WHATSAPP_GROUP_S  = "Team shyam"
REPORT_TIME       = "17:00"

# Column indices (0-based) for Apps sheet
APPS_COL_TEAM       = 1   # B
APPS_COL_COUNSELLOR = 2   # C
APPS_COL_REG_DATE   = 4   # E

# Column indices (0-based) for Admissions sheet
ADM_COL_TEAM        = 3   # D
ADM_COL_COUNSELLOR  = 2   # C
ADM_COL_ADM_DATE    = 5   # F

# ─────────────────────────────────────────────
#  GOOGLE SHEETS
# ─────────────────────────────────────────────
SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.readonly"
]

def get_sheet_data(tab_name):
    for attempt in range(3):
        try:
            creds  = Credentials.from_service_account_file(CREDENTIALS_PATH, scopes=SCOPES)
            client = gspread.authorize(creds)
            sheet  = client.open_by_key(SHEET_ID).worksheet(tab_name)
            return sheet.get_all_values()
        except Exception as e:
            print(f"[SHEET ERROR] Attempt {attempt+1}/3 — {tab_name}: {e}")
            time.sleep(5)
    return []

# ─────────────────────────────────────────────
#  EXCEL READER
# ─────────────────────────────────────────────
def read_counsellors():
    """Read Counsellors.xlsx and return dict: { 'team prashanth': [names], 'team shyam': [names] }"""
    try:
        wb     = openpyxl.load_workbook(COUNSELLORS_EXCEL, read_only=True, data_only=True)
        teams  = {}
        for sheet_name in wb.sheetnames:
            ws    = wb[sheet_name]
            names = []
            first = True
            for row in ws.iter_rows(values_only=True):
                if first:
                    first = False
                    continue  # skip header
                if row[0]:
                    # Take first word only — ignore initials like "Sahana A"
                    full = str(row[0]).strip().lower()
                    first_word = full.split()[0] if full else full
                    names.append(first_word)
            teams[sheet_name.strip().lower()] = names
            print(f"[BOT] Loaded {len(names)} counsellors from '{sheet_name}'")
        wb.close()
        return teams
    except Exception as e:
        print(f"[EXCEL ERROR] {e}")
        return {}

# ─────────────────────────────────────────────
#  DATE PARSER
# ─────────────────────────────────────────────
def parse_date(date_str):
    for fmt in ["%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"]:
        try:
            return datetime.strptime(str(date_str).strip(), fmt).date()
        except:
            continue
    return None

def is_current_month(date_str):
    parsed = parse_date(date_str)
    if not parsed:
        return False
    today = datetime.now().date()
    return parsed.month == today.month and parsed.year == today.year

# ─────────────────────────────────────────────
#  COUNT APPS AND ADMISSIONS PER COUNSELLOR
# ─────────────────────────────────────────────
def get_first_name(counsellor_str):
    """Extract first name before '/' and lowercase."""
    return counsellor_str.split("/")[0].strip().lower() if counsellor_str else ""

def count_apps_per_counsellor():
    """Count rows in All Universities applications for current month per counsellor first name."""
    rows   = get_sheet_data(TAB_APPS)
    counts = {}
    for row in rows:
        if len(row) <= APPS_COL_COUNSELLOR:
            continue
        counsellor = row[APPS_COL_COUNSELLOR].strip()
        reg_date   = row[APPS_COL_REG_DATE].strip() if len(row) > APPS_COL_REG_DATE else ""
        if not counsellor or not is_current_month(reg_date):
            continue
        first = get_first_name(counsellor)
        counts[first] = counts.get(first, 0) + 1
    return counts

def count_adm_per_counsellor():
    """Count rows in Payment done Roll no pending for current month per counsellor first name."""
    rows   = get_sheet_data(TAB_ADM)
    counts = {}
    today  = datetime.now()
    print(f"[ADM] Counting admissions for {today.strftime('%B %Y')} from '{TAB_ADM}'...")
    for row in rows:
        if len(row) <= ADM_COL_COUNSELLOR:
            continue
        counsellor = row[ADM_COL_COUNSELLOR].strip()
        adm_date   = row[ADM_COL_ADM_DATE].strip() if len(row) > ADM_COL_ADM_DATE else ""
        if not counsellor or not adm_date:
            continue
        if not is_current_month(adm_date):
            continue
        first = get_first_name(counsellor)
        counts[first] = counts.get(first, 0) + 1
        print(f"[ADM] Counted: '{counsellor}' → first='{first}' | adm_date='{adm_date}'")
    return counts

# ─────────────────────────────────────────────
#  BUILD REPORT
# ─────────────────────────────────────────────
def build_report(teams, apps_counts, adm_counts, team_filter=None):
    """Build report. team_filter: 'team prashanth' or 'team shyam' — shows only that team's scorecard."""
    now        = datetime.now()
    month_str  = now.strftime("%B %Y")
    date_str   = now.strftime("%d %B %Y")
    line       = "\u2501" * 24
    medals     = ["\U0001f947", "\U0001f948", "\U0001f949"]

    msg = (
        f"\U0001f4ca DAILY COUNSELLOR REPORT\n"
        f"{line}\n"
        f"\U0001f4c5 {month_str}  |  As of {date_str}\n"
        f"{line}\n"
    )

    team_display = {
        "team prashanth": "TEAM PRASHANTH",
        "team shyam":     "TEAM SHYAM",
    }

    team_totals = {}

    for team_key, display_name in team_display.items():
        # Only show this team's scorecard if filter is set
        if team_filter and team_key != team_filter:
            # Still need totals — collect counts but don't add to msg
            counsellors = teams.get(team_key, [])
            total_apps = sum(apps_counts.get(n.strip().lower().split()[0], 0) for n in counsellors)
            total_adm  = sum(adm_counts.get(n.strip().lower().split()[0], 0) for n in counsellors)
            team_totals[display_name] = (total_apps, total_adm)
            continue
        counsellors = teams.get(team_key, [])
        if not counsellors:
            continue

        # Build rows with counts
        rows = []
        for name in counsellors:
            apps = apps_counts.get(name, 0)
            adm  = adm_counts.get(name, 0)
            rows.append((name.title(), apps, adm))

        # Sort: admissions first, then applications as tiebreaker
        rows.sort(key=lambda x: (x[2], x[1]), reverse=True)

        msg += f"\u2b50 *{display_name}*\n{line}\n"

        total_apps = 0
        total_adm  = 0
        for i, (name, apps, adm) in enumerate(rows):
            prefix = medals[i] if i < 3 else f"{i+1}."
            padded = name.ljust(14)
            msg += f"{prefix} {padded}\u2192  {apps} Apps  |  {adm} Adm\n"
            total_apps += apps
            total_adm  += adm

        team_totals[display_name] = (total_apps, total_adm)
        msg += f"{line}\n"

    # Team totals comparison
    msg += f"\U0001f3c6 *TEAM TOTALS*\n{line}\n"

    # Sort teams by apps first for display, admissions as tiebreaker
    sorted_teams = sorted(team_totals.items(), key=lambda x: (x[1][0], x[1][1]), reverse=True)

    for i, (tname, (tapps, tadm)) in enumerate(sorted_teams):
        padded = tname.ljust(16)
        msg += f"{padded}\u2192  {tapps} Apps  |  {tadm} Adm\n"

    # Apps and admissions leader lines
    if len(sorted_teams) >= 2:
        t1_name, (t1_apps, t1_adm) = sorted_teams[0]
        t2_name, (t2_apps, t2_adm) = sorted_teams[1]

        # Applications leader
        if t1_apps > t2_apps:
            apps_leader = t1_name
        elif t2_apps > t1_apps:
            apps_leader = t2_name
        else:
            apps_leader = None

        # Admissions leader
        if t1_adm > t2_adm:
            adm_leader = t1_name
        elif t2_adm > t1_adm:
            adm_leader = t2_name
        else:
            adm_leader = None

        # Both same team
        if apps_leader and adm_leader and apps_leader == adm_leader:
            msg += f"\U0001f3c6 {apps_leader} leads on both!\n"
        else:
            if apps_leader:
                msg += f"\U0001f4cb {apps_leader} leads on applications!\n"
            else:
                msg += f"\U0001f91d Both teams tied on applications!\n"
            if adm_leader:
                msg += f"\U0001f3c6 {adm_leader} leads on admissions!\n"
            else:
                msg += f"\U0001f91d Both teams tied on admissions!\n"

    msg += f"{line}\n"
    quote = get_report_quote()
    msg += f'\U0001f4ac \"{quote}\"'
    return msg.strip()

# ─────────────────────────────────────────────
#  SELENIUM / WHATSAPP
# ─────────────────────────────────────────────
chrome_options = Options()
chrome_options.add_argument(f"user-data-dir={CHROME_PROFILE}")
chrome_options.add_argument("--start-maximized")
chrome_options.binary_location = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
chrome_options.add_experimental_option("detach", True)

driver = None
wait   = None

def init_driver():
    global driver, wait
    service = Service(ChromeDriverManager().install())
    driver  = webdriver.Chrome(service=service, options=chrome_options)
    wait    = WebDriverWait(driver, 60)
    driver.get("https://web.whatsapp.com")
    print("[BOT] Waiting for WhatsApp Web — scan QR if needed...")
    wait.until(EC.presence_of_element_located((By.XPATH, '//span[@title]')))
    time.sleep(5)
    print("[BOT] WhatsApp loaded.")

def open_chat(name):
    try:
        # Try exact match first
        chats = driver.find_elements(By.XPATH, f'//span[@title="{name}"]')
        if chats:
            chats[0].click()
            time.sleep(2)
            return True

        # Search by typing in search box
        search_box = wait.until(EC.presence_of_element_located(
            (By.XPATH, '//div[@contenteditable="true"][@data-tab="3"]')
        ))
        search_box.click()
        time.sleep(0.5)
        search_box.send_keys(Keys.CONTROL, 'a')
        search_box.send_keys(Keys.DELETE)
        time.sleep(0.5)
        search_box.send_keys(name)
        time.sleep(3)

        # Try exact match after search
        results = driver.find_elements(By.XPATH, f'//span[@title="{name}"]')
        if results:
            results[0].click()
            time.sleep(2)
            return True

        # Try partial match
        results = driver.find_elements(By.XPATH, f'//span[contains(@title, "{name}")]')
        if results:
            results[0].click()
            time.sleep(2)
            return True

        # Click first result in list
        results = driver.find_elements(By.XPATH, '//span[@data-testid="cell-frame-title"]')
        if results:
            results[0].click()
            time.sleep(2)
            return True

        print(f"[CHAT ERROR] Could not find '{name}' — skipping.")
        return False
    except Exception as e:
        print(f"[CHAT ERROR] Could not open '{name}': {e}")
        return False

def send_message(text):
    try:
        import pyperclip
        pyperclip.copy(text)
        text_box = wait.until(EC.presence_of_element_located(
            (By.XPATH, '//div[@contenteditable="true"][@data-tab="10"]')
        ))
        text_box.click()
        time.sleep(0.5)
        webdriver.ActionChains(driver).key_down(Keys.CONTROL).send_keys('v').key_up(Keys.CONTROL).perform()
        time.sleep(1)
        text_box = driver.find_element(By.XPATH, '//div[@contenteditable="true"][@data-tab="10"]')
        text_box.send_keys(Keys.ENTER)
        time.sleep(2)
        print("[BOT] Message sent.")
    except Exception as e:
        print(f"[SEND ERROR] {e}")

# ─────────────────────────────────────────────
#  MAIN REPORT FUNCTION
# ─────────────────────────────────────────────
def post_counsellor_report():
    print(f"[BOT] Building counsellor report — {datetime.now().strftime('%d %b %Y %I:%M %p')}")

    # 1. Read Excel fresh every time
    teams = read_counsellors()
    if not teams:
        print("[BOT] No counsellors found in Excel — skipping report.")
        return

    # 2. Count apps and admissions
    apps_counts = count_apps_per_counsellor()
    adm_counts  = count_adm_per_counsellor()

    print(f"[BOT] Apps counts: {apps_counts}")
    print(f"[BOT] Adm counts: {adm_counts}")

    # 3. Build and send — Team Prashanth group gets Prashanth scorecard
    msg_p = build_report(teams, apps_counts, adm_counts, team_filter="team prashanth")
    if open_chat(WHATSAPP_GROUP_P):
        send_message(msg_p)
        print("[BOT] Prashanth report posted.")
    else:
        print(f"[BOT] Could not open '{WHATSAPP_GROUP_P}'.")

    time.sleep(5)

    # 4. Team Shyam group gets Shyam scorecard
    print(f"[BOT] Now opening Team Shyam group: '{WHATSAPP_GROUP_S}'")
    msg_s = build_report(teams, apps_counts, adm_counts, team_filter="team shyam")
    if open_chat(WHATSAPP_GROUP_S):
        send_message(msg_s)
        print("[BOT] Shyam report posted.")
    else:
        print(f"[BOT] Could not open '{WHATSAPP_GROUP_S}'.")

# ─────────────────────────────────────────────
#  SCHEDULER
# ─────────────────────────────────────────────
def setup_schedule():
    schedule.every().day.at(REPORT_TIME).do(post_counsellor_report)
    print(f"[BOT] Scheduler ready — report posts daily at {REPORT_TIME}")

# ─────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────
if __name__ == "__main__":
    print("=" * 52)
    print("  📊 COUNSELLOR REPORT BOT")
    print("  Phoenix Online Education — Floor 2")
    print("=" * 52)

    init_driver()
    setup_schedule()

    print("\n[BOT] Running. Press Ctrl+C to stop.")
    print(f"[BOT] Report will post daily at {REPORT_TIME} to '{WHATSAPP_GROUP_P}' and '{WHATSAPP_GROUP_S}'.\n")

    while True:
        schedule.run_pending()
        time.sleep(1)
